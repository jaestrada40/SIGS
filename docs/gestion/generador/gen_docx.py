import json, re
import os
DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..') + '/'   # carpeta docs/gestion
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from mdtab import tables

G = DIR
OUT = DIR + 'SISG_Grupo4_Plan_de_gestion_IEEE.docx'
P = json.load(open('plan.json')); E = P['evm']; C = P['cadena_bd2']; RES = json.load(open('reserva.json'))['res_total']
T = {f: tables(G + f) for f in ['01_acta_y_alcance.md', '02_edt_cronograma_costos.md', '03_interesados_comunicaciones.md', '04_riesgos_calidad.md', '05_apendice_declaracion_ia.md']}
t01, t02, t03, t04, t05 = (T[k] for k in T)
FONT = 'Times New Roman'
FIG = 'figuras/'

doc = Document()
# ---- estilos base (IEEE: Times New Roman, interlineado sencillo, texto justificado) ----
def style_font(st, size, bold=None, italic=None, color=RGBColor(0, 0, 0)):
    st.font.name = FONT; st.font.size = Pt(size); st.font.color.rgb = color
    if bold is not None: st.font.bold = bold
    if italic is not None: st.font.italic = italic
    rpr = st.element.get_or_add_rPr(); rf = rpr.find(qn('w:rFonts'))
    if rf is None: rf = OxmlElement('w:rFonts'); rpr.append(rf)
    for a in ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs'): rf.set(qn(a), FONT)
    for a in ('w:asciiTheme', 'w:hAnsiTheme', 'w:eastAsiaTheme', 'w:cstheme'):
        if rf.get(qn(a)) is not None: del rf.attrib[qn(a)]
n = doc.styles['Normal']; style_font(n, 10); n.paragraph_format.space_after = Pt(0); n.paragraph_format.space_before = Pt(0)
n.paragraph_format.line_spacing = 1.0; n.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
h1 = doc.styles['Heading 1']; style_font(h1, 10, bold=False); h1.font.small_caps = True; h1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
h1.paragraph_format.space_before = Pt(10); h1.paragraph_format.space_after = Pt(4); h1.paragraph_format.keep_with_next = True
h2 = doc.styles['Heading 2']; style_font(h2, 10, bold=False, italic=True); h2.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
h2.paragraph_format.space_before = Pt(6); h2.paragraph_format.space_after = Pt(2); h2.paragraph_format.keep_with_next = True
lb = doc.styles['List Bullet']; style_font(lb, 10); lb.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

# ---- página: carta, márgenes IEEE ----
s0 = doc.sections[0]
s0.page_width, s0.page_height = Cm(21.59), Cm(27.94)
s0.top_margin, s0.bottom_margin, s0.left_margin, s0.right_margin = Cm(1.9), Cm(2.54), Cm(1.57), Cm(1.57)
FULL = 21.59 - 2 * 1.57; COL = 8.89

def set_cols(sec, n):
    sp = sec._sectPr; c = sp.find(qn('w:cols'))
    if c is None: c = OxmlElement('w:cols'); sp.append(c)
    c.set(qn('w:num'), str(n)); c.set(qn('w:space'), '357')
    if n == 1 and c.get(qn('w:equalWidth')) is not None: del c.attrib[qn('w:equalWidth')]
mode = {'cols': 1}
def want(n):
    if mode['cols'] != n:
        sec = doc.add_section(WD_SECTION.CONTINUOUS); set_cols(sec, n); mode['cols'] = n

# ---- helpers de texto ----
INL = re.compile(r'(\*\*.+?\*\*|`.+?`|\*[^*\s][^*]*?\*)')
def fmt_run(p, text, size=10, bold=False, italic=False, small=False, code=False, color=None):
    r = p.add_run(text); f = r.font
    f.name = 'Courier New' if code else FONT; f.size = Pt(size - (1 if code else 0)); f.bold = bold; f.italic = italic; f.small_caps = small
    rpr = r._element.get_or_add_rPr(); rf = rpr.find(qn('w:rFonts'))
    if rf is None: rf = OxmlElement('w:rFonts'); rpr.insert(0, rf)
    for a in ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs'): rf.set(qn(a), 'Courier New' if code else FONT)
    if color: f.color.rgb = color
    return r
def rich(p, text, size=10, base_bold=False, base_italic=False):
    for part in INL.split(text):
        if not part: continue
        if part.startswith('**') and part.endswith('**'): fmt_run(p, part[2:-2], size, True, base_italic)
        elif part.startswith('`') and part.endswith('`'): fmt_run(p, part[1:-1], size, base_bold, base_italic, code=True)
        elif part.startswith('*') and part.endswith('*') and len(part) > 2: fmt_run(p, part[1:-1], size, base_bold, True)
        else: fmt_run(p, part, size, base_bold, base_italic)
def para(text, size=10, align=None, indent=True, after=0, keep=False):
    p = doc.add_paragraph(); rich(p, text, size)
    if align is not None: p.alignment = align
    if indent: p.paragraph_format.first_line_indent = Cm(0.4)
    p.paragraph_format.space_after = Pt(after)
    if keep: p.paragraph_format.keep_with_next = True
    return p
def bullets(items):
    for it in items:
        p = doc.add_paragraph(style='List Bullet'); rich(p, it)
def roman(n):
    out = ''
    for v, s in ((10, 'X'), (9, 'IX'), (5, 'V'), (4, 'IV'), (1, 'I')):
        while n >= v: out += s; n -= v
    return out
ROM = [roman(i) for i in range(1, 60)]
cnt = {'sec': 0, 'sub': 0, 'tab': 0, 'fig': 0}
def H1(text):
    want(2); cnt['sec'] += 1; cnt['sub'] = 0
    p = doc.add_paragraph(style='Heading 1'); fmt_run(p, f"{ROM[cnt['sec'] - 1]}. {text.upper()}", 10, small=False)
def H2(text):
    want(2); cnt['sub'] += 1
    p = doc.add_paragraph(style='Heading 2'); fmt_run(p, f"{chr(64 + cnt['sub'])}. {text}", 10, italic=True)
def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); sh = OxmlElement('w:shd'); sh.set(qn('w:val'), 'clear'); sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), fill); tcPr.append(sh)
def cell_margins(tbl, tw=60):
    tblPr = tbl._tbl.tblPr; m = OxmlElement('w:tblCellMar')
    for side in ('top', 'left', 'bottom', 'right'):
        e = OxmlElement(f'w:{side}'); e.set(qn('w:w'), str(tw)); e.set(qn('w:type'), 'dxa'); m.append(e)
    look = tblPr.find(qn('w:tblLook'))
    (look.addprevious(m) if look is not None else tblPr.append(m))
def table(caption, t, cols=None, widths=None, fs=8, width=FULL):
    """Tabla IEEE: título arriba, centrado, en versalitas; ocupa todo el ancho de la página."""
    want(1); cnt['tab'] += 1
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(6); p.paragraph_format.keep_with_next = True
    fmt_run(p, f"TABLA {ROM[cnt['tab'] - 1]}", 8, small=True); p.add_run().add_break()
    fmt_run(p, caption.upper(), 8, small=True)
    hdr = t['header']; rows = t['rows']
    keep = cols if cols is not None else list(range(len(hdr)))
    tb = doc.add_table(rows=1 + len(rows), cols=len(keep)); tb.style = 'Table Grid'; tb.alignment = WD_TABLE_ALIGNMENT.CENTER; tb.autofit = False
    cell_margins(tb)
    w = widths or [max(4, min(36, max(len(re.sub(r'[*`]', '', str(x[k]))) for x in ([hdr] + rows)))) for k in keep]
    tot = sum(w); cw = [Cm(width * x / tot) for x in w]
    for ci, k in enumerate(keep):
        c = tb.rows[0].cells[ci]; c.width = cw[ci]; shade(c, 'D9D9D9'); c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
        rich(c.paragraphs[0], hdr[k], fs, base_bold=True)
        for r in c.paragraphs[0].runs: r.font.bold = True
    for ri, row in enumerate(rows, 1):
        for ci, k in enumerate(keep):
            c = tb.rows[ri].cells[ci]; c.width = cw[ci]; c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT; rich(c.paragraphs[0], row[k] if k < len(row) else '', fs)
    tr = tb.rows[0]._tr; trPr = tr.get_or_add_trPr(); th = OxmlElement('w:tblHeader'); th.set(qn('w:val'), 'true'); trPr.append(th)
    for row in tb.rows:
        trPr = row._tr.get_or_add_trPr(); cs = OxmlElement('w:cantSplit'); cs.set(qn('w:val'), 'true'); trPr.append(cs)
    sp = doc.add_paragraph(); sp.paragraph_format.space_after = Pt(4)
def figure(path, caption, width_cm):
    want(1); cnt['fig'] += 1
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(6); p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(FIG + path, width=Cm(width_cm))
    c = doc.add_paragraph(); c.alignment = WD_ALIGN_PARAGRAPH.CENTER; c.paragraph_format.space_after = Pt(6)
    fmt_run(c, f"Fig. {cnt['fig']}. {caption}", 8)

# =============================== CARÁTULA ===============================
def center(text, size, bold=False, italic=False, before=0, after=0):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(before); p.paragraph_format.space_after = Pt(after)
    fmt_run(p, text, size, bold, italic); return p
center('UNIVERSIDAD SAN PABLO DE GUATEMALA', 16, True, before=40)
center('Facultad de Ciencias Empresariales', 13, after=2)
center('Escuela de Ingeniería', 13, after=2)
center('Ingeniería en Sistemas y Ciencias de la Computación', 13, after=40)
center('Administración de Proyectos Tecnológicos', 15, True, after=2)
center('Código 3CE0101 · Segundo Semestre 2026 · Sección A', 12, after=40)
center('Plan de gestión del proyecto', 22, True, after=4)
center('SISG – Módulo del Grupo 4', 22, True, after=6)
center('Acta de constitución, alcance, cronograma, costos, interesados, riesgos y calidad', 13, italic=True, after=50)
center('Docente: Ing. Rafael G. Barrios Parra', 12, after=18)
center('Integrantes del Grupo 4', 12, True, after=4)
for nm in ['Javier Augusto Estrada Gordillo', 'Jonathan Eduardo Jolón García', 'Jonathan Javier Soberanis Castillo']:
    center(f'{nm}     Carné: ______________', 12, after=3)
center('Guatemala, 24 de octubre de 2026', 12, before=40, after=4)
center('Borrador v0.2 · incluye declaración de uso de IA (Apéndice A)', 10, italic=True)
set_cols(doc.sections[0], 1)
# pie de la carátula vacío
doc.sections[0].footer.is_linked_to_previous = False

# =============================== TÍTULO IEEE ===============================
sec = doc.add_section(WD_SECTION.NEW_PAGE); set_cols(sec, 1)
sec.footer.is_linked_to_previous = False
fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
for kind, txt in (('begin', None), ('instr', ' PAGE '), ('end', None)):
    r = fp.add_run(); r.font.name = FONT; r.font.size = Pt(9)
    if kind == 'instr': e = OxmlElement('w:instrText'); e.set(qn('xml:space'), 'preserve'); e.text = txt
    else: e = OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'), kind)
    r._r.append(e)
pt = doc.add_paragraph(); pt.alignment = WD_ALIGN_PARAGRAPH.CENTER; pt.paragraph_format.space_after = Pt(8)
fmt_run(pt, 'Plan de gestión del proyecto SISG – Módulo del Grupo 4', 24)
pa = doc.add_paragraph(); pa.alignment = WD_ALIGN_PARAGRAPH.CENTER; pa.paragraph_format.space_after = Pt(2)
fmt_run(pa, 'Javier Augusto Estrada Gordillo, Jonathan Eduardo Jolón García y Jonathan Javier Soberanis Castillo', 11)
pf = doc.add_paragraph(); pf.alignment = WD_ALIGN_PARAGRAPH.CENTER; pf.paragraph_format.space_after = Pt(10)
fmt_run(pf, 'Escuela de Ingeniería en Sistemas y Ciencias de la Computación, Universidad San Pablo de Guatemala, Guatemala', 10, italic=True)
mode['cols'] = 1

# =============================== RESUMEN ===============================
want(2)
pr = doc.add_paragraph(); pr.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY; pr.paragraph_format.first_line_indent = Cm(0)
fmt_run(pr, 'Resumen—', 9, True, True)
rich(pr, (f"Este documento presenta el plan de gestión del módulo del Grupo 4 del Sistema Integral de Salud de Guatemala (SISG): acta de constitución, alcance, "
          f"estructura de desglose del trabajo, cronograma, costos, interesados, comunicaciones, riesgos y calidad. Con corte al 24 de octubre de 2026, el avance estimado es de "
          f"{E['avance']}% y el índice de desempeño del cronograma (SPI) es {E['SPI']}; ambos dependen de estimaciones que el grupo debe validar. La cadena de entrega de Bases de Datos II "
          f"tiene {C['holgura_dias']} días de holgura y una probabilidad aproximada de {C['prob_cumplir_pct']}% de cumplir el plazo, por lo que se identifica como el riesgo crítico del proyecto."), 9, False, True)
pk = doc.add_paragraph(); pk.paragraph_format.space_after = Pt(4); pk.paragraph_format.first_line_indent = Cm(0)
fmt_run(pk, 'Palabras clave—', 9, True, True); fmt_run(pk, 'gestión de proyectos, PMBOK 8, almacén de datos, valor ganado, salud digital, interoperabilidad', 9, False, True)

# =============================== I. INTRODUCCIÓN ===============================
H1('Introducción')
para("El Sistema Integral de Salud de Guatemala (SISG) busca que los actores del ecosistema de salud compartan datos de forma segura y consistente. El marco de referencia es el documento de la OPS [1], que propone partir de una línea base, definir una arquitectura objetivo y planear transiciones, con la interoperabilidad y la seguridad de la información como principios rectores.")
para("El Grupo 4 representa a los centros de diagnóstico y a los organismos no lucrativos, que alimentan y consumen datos. Su módulo registra pacientes, estudios y campañas de salud, y los deja listos para análisis en un almacén de datos con modelo estrella.")
para("Este plan sigue la guía del PMBOK 8 [2] y las prácticas de ISO 21502 [3], y adopta las plantillas y métodos vistos en las sesiones del curso [5]: acta de constitución, registro de interesados, plan de comunicaciones, método de precedencias, valor ganado y matriz de riesgos 3×3. La calidad del producto se especifica con ISO/IEC 25010 [4].")

# =============================== II. ACTA ===============================
H1('Acta de constitución')
H2('Título y propósito')
para("**Título:** SISG – Módulo del Grupo 4 (centros de diagnóstico y organismos no lucrativos). **Propósito:** registrar pacientes, estudios diagnósticos y campañas de salud, y dejar esos datos listos para el análisis en un almacén de datos con modelo estrella.")
H2('Justificación y caso de negocio')
para("Los centros y organismos hoy trabajan con archivos de Excel y bases aisladas. Sin un modelo común, ministerios, aseguradoras y universidades no pueden comparar estudios ni campañas por mes, tipo o lugar. El módulo entrega ese modelo común y demuestra, con datos de prueba, que los indicadores se calculan sin perder filas ni montos (reconciliación en cero).")
H2('Objetivos y criterios de éxito')
table('Objetivos SMART del proyecto', t01[0], widths=[3, 30])
bullets([
 "La reconciliación fuente–almacén de datos da diferencia 0 en filas y montos después de cada carga.",
 "Las pruebas e2e de la aplicación pasan al 100% (49 de 49 al 24 de octubre).",
 "El almacén de datos contiene 0 columnas con nombre, DPI o teléfono de pacientes.",
 "Todas las entregas incluyen carátula, formato IEEE y el apéndice de declaración de uso de IA.",
 f"El plazo de Bases de Datos II se cumple con al menos 2 días de holgura (hoy: {C['holgura_dias']} días y ≈ {C['prob_cumplir_pct']}% de probabilidad).",
 "Al menos 80% de las personas de una prueba de usabilidad registran un paciente sin ayuda."])
H2('Alcance de alto nivel y exclusiones')
para("Incluye la aplicación web con carga masiva por CSV y bitácora, el esquema fuente, el almacén de datos con su ETL y consultas, y los documentos de gestión. Excluye la integración real con los demás grupos y con el Gestor de plataforma (Grupo 8), la puesta en producción, la recuperación de contraseña, el ETL incremental, los archivos .xlsx (se acepta CSV) y los tableros de BI externos.")
H2('Interesados clave')
para("Docentes, equipo del Grupo 4, Gestor de plataforma, Ministerio de Salud/INE/Estado, hospitales, seguro social, centros y organismos usuarios, consumidores de datos y pacientes. El registro completo se presenta en la sección V.")
H2('Hitos y cronograma preliminar')
table('Hitos preliminares', t01[1], widths=[6, 30])
H2('Presupuesto estimado')
para(f"No hay gasto monetario directo (Q 0): se usan Docker y software libre en equipos propios. El costo se mide en horas-persona: BAC = {E['BAC']} h; reserva de contingencia {E['contingencia']} h ({E['contingencia_pct']}%); reserva de gestión {E['gestion_5pct']} h (5%); presupuesto con reservas = {round(E['BAC'] + E['contingencia'] + E['gestion_5pct'], 1)} h. Las horas son estimaciones iniciales por validar.")
H2('Riesgos y supuestos')
table('Riesgos principales del acta', t01[2], widths=[3, 30, 6])
para("**Supuestos:** el grupo dispone de Docker y de un equipo para ejecutar las bases; los datos de prueba son sintéticos; los demás grupos entregarán datos que el ETL pueda leer. **Restricciones:** plazo fijo del 28 de noviembre, equipo de 3 personas, y un curso que enseña SQL Server frente a un prototipo en MySQL.")
H2('Director de proyecto y su autoridad')
para("**Director de proyecto (líder):** por asignar por el grupo. **Autoridad:** aprobar cambios de alcance dentro del equipo y registrarlos; reasignar paquetes de trabajo; usar las reservas; convocar la reunión semanal; y escalar a los docentes lo que excede la autoridad del equipo, como el motor de base de datos o el alcance de una entrega.")
table('Roles propuestos', t01[3], widths=[10, 22, 8])
H2('Firma del patrocinador')
para("Patrocinadores académicos: Ing. Rafael G. Barrios Parra (Administración de Proyectos Tecnológicos) e Inga. Claudia Alegre (Bases de Datos II). Firma y fecha: ____________________ · ____________________")

# =============================== III. ALCANCE ===============================
H1('Alcance del proyecto')
H2('Alcance del producto y del proyecto')
para("El **alcance del producto** (el qué) es una aplicación web con autenticación, módulos de pacientes, estudios y campañas, carga masiva y bitácora; un esquema fuente en tercera forma normal; y un almacén de datos con ETL y consultas analíticas. El **alcance del proyecto** (el cómo y el cuánto trabajo) incluye la documentación de gestión, el desarrollo y las pruebas del producto, la exposición final y el informe de cierre.")
H2('Entregables y criterios de aceptación')
table('Criterios de aceptación por entregable', t01[4], widths=[10, 30])
H2('Enfoque híbrido e historias de usuario')
para("Los entregables de Bases de Datos II tienen contenido y fecha fijos (enfoque predictivo); la aplicación evoluciona con un backlog priorizado (enfoque ágil). Las seis historias iniciales están implementadas al 24 de octubre. La carga masiva acepta CSV, no archivos .xlsx.")
table('Historias de usuario del backlog', t01[5], widths=[3, 22, 22, 6])

# =============================== IV. EDT, CRONOGRAMA, COSTOS ===============================
H1('EDT, cronograma y costos')
H2('Estructura de desglose del trabajo')
para("La EDT descompone el 100% del alcance en cinco entregables y 20 paquetes de trabajo (Fig. 1 y Tabla del diccionario). Las horas son estimaciones iniciales; el grupo debe reemplazarlas por sus horas reales.")
figure('edt.png', 'Estructura de desglose del trabajo con horas estimadas por paquete.', 17.5)
table('Diccionario de paquetes de trabajo', t02[0], widths=[5, 26, 5, 14, 6, 9])
H2('Secuenciación por precedencia')
para("Las actividades restantes se secuencian con el método de diagramación por precedencia desde el lunes 26 de octubre (Fig. 2). Se aplican dos restricciones reales: la exposición solo puede darse el 20 y 21 de noviembre (semana 19) y la entrega de Bases de Datos II vence el 28 de noviembre a las 8:00. La revisión por pares se solapa 2 días con la redacción del documento técnico.")
figure('pdm.png', 'Diagrama de red con duración y holgura de cada actividad.', 17.5)
table('Actividades, predecesoras y holgura (días desde el 26 de octubre)', t02[1], widths=[3, 26, 6, 8, 4, 4, 4, 4, 6, 12])
H2('Ruta crítica y estimación PERT')
para(f"La exposición (J) tiene holgura 0 porque su fecha está fija por el calendario. La cadena de Bases de Datos II (B, C, D, E, K, L) termina el 25 de noviembre, con {C['holgura_dias']} días de holgura frente al límite del 28 a las 8:00. Con la estimación de tres valores, la duración esperada de la cadena es de {C['esperado_dias']} días (σ = {C['sigma']}) contra {C['limite_dias']} disponibles: z = {C['z']}, es decir, una probabilidad de cumplir de ≈ {C['prob_cumplir_pct']}%. Los valores optimista, probable y pesimista son estimaciones iniciales del borrador.")
table('Estimación PERT de la cadena de Bases de Datos II', t02[2], widths=[3, 26, 4, 4, 4, 8, 4])
H2('Cronograma')
figure('gantt.png', 'Cronograma de las actividades restantes; las líneas punteadas indican la holgura.', 17.5)
H2('Costos y valor ganado')
para(f"El costo se mide en horas-persona. BAC = {E['BAC']} h. La reserva de contingencia se fija en {E['contingencia_pct']}% ({E['contingencia']} h) porque las respuestas a los riesgos suman {RES} h y el 10% sugerido en clase ({E['contingencia_10pct']} h) no alcanza; la reserva de gestión es de {E['gestion_5pct']} h (5%). Al 24 de octubre, EV = {E['EV']} h, PV = {E['PV']} h y SPI = {E['SPI']}. El costo real (AC) aún no está registrado, por lo que CPI y EAC quedan pendientes.")
table('Indicadores de valor ganado al 24 de octubre de 2026', t02[3], widths=[9, 6, 26])
H2('Enfoque híbrido: Fase 0 y sprints')
para("El plan combina un marco predictivo (Fase 0 y hitos fin-a-inicio en la ruta crítica) con ejecución ágil por sprints, cada uno con su objetivo (Sprint Goal), su incremento y su definición de hecho [6]. El modelo de datos verificado va antes que las consultas y cada sprint construye sobre lo que el anterior dejó verificado; si un hito se atrasa, se mueve la fecha de los primeros pasos del siguiente.")
table('Etapas, objetivo y definición de hecho', t02[4], widths=[10, 12, 6, 22, 10])

# =============================== V. INTERESADOS Y COMUNICACIONES ===============================
H1('Interesados y comunicaciones')
H2('Registro de interesados')
para("Los actores salen de la tabla de grupos del proyecto integral. Poder e interés (1 a 5) son una valoración inicial del grupo; «alto» es 4 o 5.")
table('Registro de interesados', t03[0], widths=[5, 22, 6, 28, 5, 5, 12])
H2('Matriz de poder e interés y estrategias')
figure('poder.png', 'Matriz de poder e interés; los números corresponden a la Tabla de estrategias.', 12)
table('Estrategia por interesado', t03[1], widths=[20, 24, 12, 34])
H2('Plan de comunicaciones')
table('Plan de comunicaciones', t03[2], widths=[12, 12, 14, 12, 15, 9, 15])
H2('Reporte de estado')
para("El reporte sigue la plantilla del curso. CPI y EAC están pendientes hasta registrar el costo real. El reporte es un borrador generado con IA y debe ser verificado por una persona del grupo antes de enviarse.")
table('Reporte de estado al 24 de octubre de 2026', t03[3], widths=[8, 40])

# =============================== VI. RIESGOS ===============================
H1('Gestión de riesgos')
H2('Proceso y análisis cualitativo')
para("El proceso sigue el ciclo común a Watt [8], ISO 21502 [3], PMBOK 8 [2] y NIST SP 800-30 [7]: preparar, conducir, comunicar y mantener. El propósito es proteger el plazo, la calidad de las entregas y los datos de salud que manejará la aplicación; el horizonte es el 28 de noviembre de 2026 para los riesgos del proyecto y de 6 meses para los de seguridad. Cada riesgo se evalúa con la matriz 3×3 de probabilidad e impacto (niveles bajo, medio, alto y crítico) y con el porqué de cada valor, para que otra persona llegue al mismo resultado. Las estrategias para amenazas son evitar, transferir, mitigar, aceptar y escalar. La severidad de una vulnerabilidad no equivale al riesgo: una puntuación de 10.0 puede dejar un riesgo residual bajo si el equipo ya está corregido. El riesgo R1 (que la docente exigiera SQL Server) se cerró el 24 de octubre porque el grupo decidió que MySQL es definitivo.")
table('Matriz de probabilidad e impacto', t04[0], widths=[8, 10, 10, 10])
H2('Estructura de desglose de riesgos')
table('Riesgos por categoría', t04[1], widths=[8, 8, 32])
H2('Registro de riesgos')
table('Registro de riesgos (amenazas)', t04[2], widths=[3, 6, 18, 4, 4, 5, 20, 22, 4, 4])
H2('Modelo de riesgo de NIST')
para("Los riesgos de seguridad se describen con el modelo de NIST [7]: una fuente de amenaza inicia un evento que explota una vulnerabilidad y produce un impacto adverso; la condición predisponente sube o baja la probabilidad.")
table('Fuente, evento, vulnerabilidad e impacto de los riesgos de seguridad', t04[3], widths=[3, 9, 10, 14, 10, 8, 10])
H2('Plan de contingencia')
para("El plan de contingencia es el método alternativo que se aplica cuando un riesgo identificado ocurre; es distinto de los fondos de contingencia descritos más abajo. Si el sistema cae, los casos se registran en la plantilla CSV y se cargan después con la carga masiva.")
table('Plan de contingencia', t04[4], widths=[4, 16, 28, 6])
H2('Evidencia del laboratorio A1-b')
para("Cada integrante inventaría su equipo real y registra una fila de riesgo por sistema presente más una vulnerabilidad vigente de su versión exacta, consultada en el NVD [9] y en el aviso del fabricante [10]. Se consolida sin nombres ni números de serie. La base de la clase es CVE-2025-24054 (Windows NTLM) y CVE-2025-43300 (macOS ImageIO). Las filas marcadas por completar las llena cada integrante.")
table('Inventario y estado de los equipos del grupo', t04[5], widths=[5, 9, 6, 14, 12, 10])
H2('Oportunidades y reservas')
table('Oportunidades', t04[6], widths=[3, 28, 8, 28])
para(f"Las respuestas a las amenazas suman {RES} h, por lo que la reserva de contingencia se fija en {E['contingencia']} h ({E['contingencia_pct']}% del BAC de {E['BAC']} h). La reserva de gestión (imprevistos no identificados, bajo control del líder) es de {E['gestion_5pct']} h. La reserva se reajusta cuando cambia la priorización de riesgos.")

# =============================== VII. CALIDAD ===============================
H1('Plan de gestión de la calidad')
H2('Especificaciones medibles')
para("La calidad se define contra requisitos documentados y se trata como principio de PMBOK 8 [2] en los siete dominios, con las prácticas de planificar, asegurar y controlar la calidad de ISO 21502 [3]. El producto se evalúa con cuatro características de ISO/IEC 25010 [4]: adecuación funcional, fiabilidad, seguridad y usabilidad. Los mínimos y metas de usabilidad y desempeño son propuestas del grupo y deben validarse antes de usarse como criterio de aceptación.")
table('Especificación de calidad', t04[7], widths=[14, 12, 16, 6, 6, 14])
H2('Calidad en los siete dominios')
table('Integración de la calidad en los dominios de PMBOK 8', t04[8], widths=[7, 18, 30])
H2('Defectos encontrados y corregidos')
table('Registro de defectos del desarrollo', t04[9], widths=[24, 16, 24])

# =============================== VIII. CONCLUSIONES ===============================
H1('Conclusiones')
bullets([
 f"Con corte al 24 de octubre, el avance estimado es de {E['avance']}%: están hechos el ER de las fuentes (entregado), el modelo estrella, el ETL, la aplicación con carga masiva, MFA, roles y auditoría, 49 pruebas e2e en verde y la reconciliación fuente–almacén con diferencia 0.",
 f"El riesgo crítico es el plazo de Bases de Datos II: {C['holgura_dias']} días de holgura y ≈ {C['prob_cumplir_pct']}% de probabilidad de cumplir, además de coincidir con el examen final.",
 "Decisión tomada: el almacén de datos es en MySQL (definitivo desde el 24 de octubre). Faltan por medir el costo real (AC) en horas para calcular CPI y EAC, y las mediciones de usabilidad y desempeño.",
 "Las horas, los porcentajes de avance, los valores PERT y las valoraciones de riesgo e interesados son estimaciones iniciales y deben validarse con el grupo antes de reportarse como resultados."])

# =============================== REFERENCIAS ===============================
want(2)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(10); p.paragraph_format.space_after = Pt(4); p.paragraph_format.keep_with_next = True
fmt_run(p, 'Referencias', 10, small=True)
REFS = [
 "Organización Panamericana de la Salud, *Bases para la construcción del Marco de Arquitectura de Salud Pública para Guatemala: Línea de base*, OPS/OMS Guatemala, 2024.",
 "Project Management Institute, *Guía de los Fundamentos para la Dirección de Proyectos (Guía del PMBOK)*, 8.ª ed., 2025.",
 "ISO, *ISO 21502:2020 – Project, programme and portfolio management – Guidance on project management*, 2020.",
 "ISO/IEC, *ISO/IEC 25010:2023 – Systems and software quality models*, 2023.",
 "R. G. Barrios Parra, *Administración de Proyectos Tecnológicos: material de las sesiones 2, 3, 5, 10, 11, 13 y de la sesión del 24 de octubre (gestión de riesgos, laboratorio A1-b y roadmap híbrido)*, Universidad San Pablo de Guatemala, 2026.",
 "K. Schwaber y J. Sutherland, *The Scrum Guide*, 2020, https://scrumguides.org/scrum-guide.html.",
 "Joint Task Force, *NIST SP 800-30 Rev. 1: Guide for Conducting Risk Assessments*, National Institute of Standards and Technology, 2012.",
 "A. Watt, *Project Management*, cap. 16, 2.ª ed., BCcampus, 2014.",
 "NIST, *National Vulnerability Database (NVD)*, https://nvd.nist.gov, consultado el 24 de octubre de 2026.",
 "Apple, *About the security content of macOS Golden Gate 27*, 14 de septiembre de 2026 (entradas actualizadas el 19 de octubre), https://support.apple.com/en-us/149035."]
for i, r in enumerate(REFS, 1):
    q = doc.add_paragraph(); q.paragraph_format.left_indent = Cm(0.8); q.paragraph_format.first_line_indent = Cm(-0.8); q.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    fmt_run(q, f"[{i}]\t", 8); rich(q, r, 8)

# =============================== APÉNDICE A ===============================
want(1)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(12); p.paragraph_format.keep_with_next = True
fmt_run(p, 'Apéndice A. Declaración de uso de IA', 10, small=True)
para("Formato del Estándar de Declaración de uso de IA para Estudiantes del curso. Los campos marcados con ⚠ solo los puede completar el grupo con honestidad: no se rellenaron porque se desconoce cuánto de este material revisó y modificó el grupo.", 9, indent=False, after=4)
table('Declaración de uso de IA', t05[0], widths=[8, 40], fs=8)
para("**Datos estimados por la IA que el grupo debe validar:** horas por paquete de trabajo, porcentajes de avance, valores optimista, probable y pesimista (PERT), poder e interés de los interesados, probabilidad e impacto de los riesgos, y los mínimos y metas de usabilidad y desempeño.", 9, indent=False)

z = doc.settings.element.find(qn('w:zoom'))
if z is not None and z.get(qn('w:percent')) is None: z.set(qn('w:percent'), '100')
doc.core_properties.title = 'Plan de gestión del proyecto SISG – Módulo del Grupo 4'
doc.core_properties.author = 'Grupo 4'; doc.core_properties.language = 'es-GT'
doc.save(OUT); print('guardado', OUT, 'tablas:', cnt['tab'], 'figuras:', cnt['fig'])
